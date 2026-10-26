"""Formats beyond email / WhatsApp / PDF, so unfamiliar uploads are read instead of skipped.

  text        notes, .md, .csv, .json, saved .html pages  -> doc_page units (~6k-char chunks)
  docx        Word documents (stdlib zip + XML, no extra dependency) -> doc_page units
  xlsx        spreadsheets that are NOT time grids -> one doc_page per sheet (cells as rows of text)
  sms_backup  Android "SMS Backup & Restore" XML -> chat_message units, one thread per phone number
"""
from __future__ import annotations
import html as htmlmod, re, zipfile, io
import xml.etree.ElementTree as ET
from datetime import datetime, timezone, timedelta
from ..models import Unit, Locator, stable_id
from ..normalize import clean_text, canonical_email
from .document_parser import classify_page

CHUNK = 6000


def _chunks(text: str, size: int = CHUNK) -> list[tuple[int, str]]:
    """Split on paragraph boundaries; returns (offset, chunk)."""
    out, start = [], 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):
            cut = text.rfind("\n\n", start + size // 2, end)
            end = cut if cut > 0 else end
        out.append((start, text[start:end]))
        start = end
    return out or [(0, "")]


def _doc_units(source, text: str, kind: str) -> list[Unit]:
    units = []
    for i, (off, chunk) in enumerate(_chunks(text)):
        dtype, _ = classify_page(chunk)
        units.append(Unit(unit_id=stable_id(source.source_id, kind, str(i)), unit_type="doc_page",
                          locator=Locator(source_id=source.source_id, kind=kind, page=i + 1,
                                          char_start=off, char_end=off + len(chunk)),
                          text=clean_text(chunk), raw_text=chunk,
                          meta={"doc_type_hint": dtype, "segment": 0, "backend": kind}))
    return units


def parse_text(store, source) -> list[Unit]:
    raw = store.raw_bytes(source.source_id)
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("utf-16", "replace")
    if re.search(r"<html|<body|<div", text[:3000], re.I):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(text, "html.parser")
        for t in soup(["script", "style", "head"]):
            t.decompose()
        text = soup.get_text("\n")
    source.meta["document"] = {"backend": "text"}
    return _doc_units(source, text, "text_chunk")


def parse_docx(store, source) -> list[Unit]:
    with zipfile.ZipFile(io.BytesIO(store.raw_bytes(source.source_id))) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    paras = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        runs = re.findall(r"<w:t(?: [^>]*)?>(.*?)</w:t>", p, re.S)
        tabs = "\t".join(runs) if "<w:tab/>" in p else "".join(runs)
        paras.append(htmlmod.unescape(tabs))
    source.meta["document"] = {"backend": "docx"}
    return _doc_units(source, "\n".join(paras), "docx_chunk")


def parse_xlsx_text(store, source) -> list[Unit]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(store.raw_bytes(source.source_id)), data_only=True, read_only=True)
    units = []
    for si, ws in enumerate(wb.worksheets):
        lines = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                lines.append(" | ".join(cells).rstrip(" |"))
            if len(lines) > 2000:
                lines.append("… (truncated)"); break
        text = f"Sheet: {ws.title}\n" + "\n".join(lines)
        dtype, _ = classify_page(text)
        units.append(Unit(unit_id=stable_id(source.source_id, "sheet", str(si)), unit_type="doc_page",
                          locator=Locator(source_id=source.source_id, kind="sheet", page=si + 1),
                          text=clean_text(text), raw_text=text,
                          meta={"doc_type_hint": dtype, "segment": si, "backend": "xlsx", "sheet": ws.title}))
    source.meta["document"] = {"backend": "xlsx", "sheets": len(units)}
    return units


def parse_sms_backup(store, source, self_aliases: dict, tz_hours: float) -> list[Unit]:
    root = ET.fromstring(store.raw_bytes(source.source_id))
    rows = []
    for el in root.iter("sms"):
        a = el.attrib
        try:
            utc = datetime.fromtimestamp(int(a.get("date", "0")) / 1000, tz=timezone.utc)
        except ValueError:
            continue
        rows.append((a.get("address", "?"), utc, a))
    rows.sort(key=lambda r: (r[0], r[1]))                 # group by thread, then time
    me = next(iter(self_aliases), "Me")
    units = []
    for i, (addr, utc, a) in enumerate(rows):
        sent = a.get("type") == "2"
        contact = a.get("contact_name") if a.get("contact_name") not in (None, "", "(Unknown)") else addr
        who = me if sent else contact
        local = (utc + timedelta(hours=tz_hours)).replace(tzinfo=None)
        body = a.get("body", "")
        flags = ["contains_otp"] if re.search(r"\b(otp|one[- ]time password|verification code)\b", body, re.I) else []
        units.append(Unit(
            unit_id=stable_id(source.source_id, "sms", str(i)), unit_type="chat_message",
            locator=Locator(source_id=source.source_id, kind="sms", line_start=i + 1, line_end=i + 1),
            text=clean_text(body), raw_text=body, timestamp_utc=utc.isoformat(),
            timestamp_raw=a.get("readable_date") or utc.isoformat(), timestamp_confidence="exact",
            author={"display": who, "address": None if sent else addr,
                    "canonical": self_aliases.get(me) if sent else addr},
            recipients=[] if not sent else [{"display": contact, "address": addr, "canonical": addr}],
            thread_key=f"sms:{addr}", meta={"local_time": local.isoformat(), "direction": "sent" if sent else "received",
                                           "thread": contact}, flags=flags))
    source.meta["whatsapp"] = {"day_first": True, "assumed_tz_hours": tz_hours, "messages": len(units),
                               "participants": sorted({(u.author or {}).get("display") for u in units}), "kind": "sms"}
    return units
